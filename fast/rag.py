from ollama import chat, embeddings, Tool
from fastapi import FastAPI, HTTPException
from fastapi import UploadFile, File, Form
from typing import List, Optional
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from qdrant_client import QdrantClient
from qdrant_client.models import PointStruct
import json
import os
import shutil
import base64

import fitz
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
from bs4 import BeautifulSoup
import csv
import re
import uuid

app = FastAPI()
qdrant = QdrantClient(
    host="localhost",
    port=6333
)
UPLOAD_FOLDER = "uploads"
DOCUMENT_FOLDER = os.path.join(
    UPLOAD_FOLDER,
    "documents"
)
IMAGE_FOLDER = os.path.join(
    UPLOAD_FOLDER,
    "images"
)

ALLOWED_IMAGES = [
    "image/png",
    "image/jpeg",
    "image/jpg"
]
ALLOWED_DOCUMENTS = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "text/plain": ".txt",
    "text/csv": ".csv",
    "text/markdown": ".md",
    "application/json": ".json",
    "text/html": ".html"
}

@app.post("/chat")
def sendChat(message: str = Form(...), images: list[UploadFile] | None = File(None), documents: list[UploadFile] | None = File(None)):
    uploadDocuments = []
    uploadImages = []
    messages = [
        {
            "role": "system",
            "content": """
                You are Kwanza AI, developed by Arsal Fahrulloh.

                You can answer general questions normally.

                Use tools only when you need:
                - search web

                Do not use tools for general knowledge, casual conversation, or conceptual explanations.

                Always use tool results when user asks for personal or searching website information.
                Do not make up database information.
            """
        },
        {
            "role": "user",
            "content": message
        }
    ]
    
    if(documents):
        for document in documents:
            if(document.content_type not in ALLOWED_DOCUMENTS):
                raise HTTPException(
                    status_code=400,
                    detail=f"{document.filename} tidak di dukung"
                )
            
            documentPath = os.path.join(
                DOCUMENT_FOLDER,
                document.filename
            )

            with open(documentPath, "wb") as buffer:
                shutil.copyfileobj(
                    document.file,
                    buffer
                )

            size = os.path.getsize(documentPath)
            uploadDocuments.append({
                "file_name": document.filename,
                "file_path": documentPath,
                "size": size,
                "file_type": document.content_type
            })

            rawText = extract_text(documentPath, document.content_type)
            text = normalize_text(rawText)
            text = clean_text(text)
            text = fix_structure(text)
            text = final_clean(text)
            chunks = chunk_text(text)
            storeQdrant(file_name=document.filename, file_path=documentPath, file_type=document.content_type, chunks=chunks)
            question = embeddings(
                model="mxbai-embed-large",
                prompt=message
            )
            
        result = qdrant.query_points(
            collection_name="documents",
            query=question.embedding,
            with_payload=True,
            limit=5,
        )
        documentContext = "\n\n".join(
            point.payload["text"]
            for point in result.points
        )
        messages.append(
            {
                "role": "system",
                "content": f"""
                    Use the provided document context to answer the user's question.

                    Rules:
                    - Use document context as the primary source of truth.
                    - Answer only based on the provided context.
                    - Do not invent or assume information.
                    - If the answer is not available in the context, reply exactly:
                    Information not found in the uploaded document.

                    Document Context:
                    {documentContext}
                    """
            }
        )

    if(images):
        for image in images:
            if(image.content_type not in ALLOWED_IMAGES):
                raise HTTPException(
                    status_code=400,
                    detail=f"{image.filename} tidak di dukung"
                )
            
            imagePath = os.path.join(
                IMAGE_FOLDER,
                image.filename
            )

            with open(imagePath, "wb") as buffer:
                shutil.copyfileobj(
                    image.file,
                    buffer
                )

            size = os.path.getsize(imagePath)
            with open(imagePath, "rb") as f:
                base64Image = base64.b64encode(f.read()).decode("utf-8")
                uploadImages.append(base64Image)

    response = chat(
        model="qwen3.5:4b",
        stream=False,
        messages=messages,
        tools=[
            {
                "type": "function",
                "function": {
                    "name": "search_web",
                    "description": "Search information from the internet",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "query": {
                                "type": "string",
                                "description": "Search query"
                            }
                        },
                        "required": ["query"]
                    }
                }
            }
        ]
    )

    toolCalls = response.message.tool_calls or []
    if(not toolCalls):
        # messages.append(response.message)
        # return messages
        return response
    
    toolMessages = []
    for tool in toolCalls:
        tool_name = tool.function.name
        argument = tool.function.arguments or {}

        result = executeTool(tool_name, arguments=argument)

        toolMessages.append({
            "role": "tool",
            "tool_name": tool_name,
            "content": json.dumps(result)
        })

    messages.append({
        "role": response.message.role,
        "content": response.message.content,
        "thinking": response.message.thinking,
        "tool_calls": response.message.tool_calls
    })
    messages.extend(toolMessages)

    finalResponse = chat(
        model="qwen3.5:4b",
        stream=False,
        messages=messages,
    )

    return finalResponse

    # messages.append({
    #     "role": finalResponse.message.role,
    #     "content": finalResponse.message.content,
    #     "thinking": finalResponse.message.thinking
    # })

    # return messages

def executeTool(name: str, arguments: dict):
    if name == "search_web":
        return toolWebSearch(arguments["query"])
    return {
        "error": "Unknown tool"
    }


def toolWebSearch(query: str):
    return {
        "query": query,
        "source": "https://www.logammulia.com/id/harga-emas-hari-ini",
        "result": {
            "title": "Harga Emas Hari Ini, di tahun 2026",
            "content": """
Berat	Harga Dasar	Harga (+Pajak PPh 0.25%)
Emas Batangan
0.5 gr	1,385,000	1,388,463
1 gr	2,670,000	2,676,675
2 gr	5,280,000	5,293,200
3 gr	7,895,000	7,914,738
5 gr	13,125,000	13,157,813
10 gr	26,195,000	26,260,488
25 gr	65,362,000	65,525,405
50 gr	130,645,000	130,971,613
100 gr	261,212,000	261,865,030
250 gr	652,765,000	654,396,913
500 gr	1,305,320,000	1,308,583,300
1000 gr	2,610,600,000	2,617,126,500
Emas Batangan Gift Series
0.5 gr	1,455,000	1,458,638
1 gr	2,820,000	2,827,050
Emas Batangan Selamat Idul Fitri
5 gr	14,098,000	14,133,245
Emas Batangan Imlek
8 gr	22,214,800	22,270,337
88 gr	241,657,600	242,261,744
Emas Batangan Batik Seri III
10 gr	27,200,000	27,268,000
20 gr	53,600,000	53,734,000


Perak Murni
Berat	Harga Dasar	Harga Sudah Termasuk PPN 11%
250 gr	10,887,500	12,085,125
500 gr	20,975,000	23,282,250
Perak Heritage
Berat	Harga Dasar	Harga Sudah Termasuk PPN 11%
31.1 gr	1,852,428	2,056,195
186.6 gr	9,993,166	11,092,414
"""
        }
    }

#qdrant point
def storeQdrant(file_name: str, file_path: str, file_type: str, chunks: list):
    points = []
    for index, chunk in enumerate(chunks):
        response = embeddings(
            model="mxbai-embed-large",
            prompt=chunk
        )
        embedding = response.embedding
        points.append(PointStruct(
            id=str(uuid.uuid4()),
            vector=embedding,
            payload={
                "file_name": file_name,
                "file_path": file_path,
                "file_type": file_type,
                "chunk_index": index + 1,
                "text": chunk
            }
        ))
        
    qdrant.upsert(
        collection_name="documents",
        points=points
    )

#chunk teks
def chunk_text(
    text: str,
    chunk_size: int = 500,
    overlap: int = 100
):
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]
        chunks.append(chunk.strip())
        start += chunk_size - overlap
    return chunks

#clean teks
def normalize_text(text: str):
    if not text:
        return ""
    text = text.replace("\t", " ")
    text = text.replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ ]{2,}", " ", text)
    text = re.sub(r"\u00A0", " ", text)
    return text.strip()

def clean_text(text: str):
    if not text:
        return ""
    text = re.sub(r"[\x00-\x08\x0B-\x1F\x7F]", "", text)
    text = re.sub(r"[•●▪►■◆]", "-", text)
    text = re.sub(r"[“”]", "\"", text)
    text = re.sub(r"[‘’]", "'", text)
    text = re.sub(r"…", "...", text)
    text = re.sub(r"­", "", text)
    text = re.sub(r"\n +", "\n", text)
    text = re.sub(r" +\n", "\n", text)
    return text.strip()

def fix_structure(text: str):
    lines = []
    for line in text.split("\n"):
        line = line.strip()
        if line == "":
            continue
        lines.append(line)
    return "\n".join(lines)

def final_clean(text: str):
    text = re.sub(r"\n{2,}", "\n", text)
    text = re.sub(r" {2,}", " ", text)
    return text.strip()

#extrak teks dari dokument
def extract_text(file_path: str, content_type: str) -> str:
    if content_type == "application/pdf":
        return extract_pdf(file_path)
    elif content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return extract_docx(file_path)
    elif content_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        return extract_pptx(file_path)
    elif content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return extract_xlsx(file_path)
    elif content_type == "text/plain":
        return extract_txt(file_path)
    elif content_type == "text/csv":
        return extract_csv(file_path)
    elif content_type == "application/json":
        return extract_json(file_path)
    elif content_type == "text/html":
        return extract_html(file_path)
    elif content_type == "text/markdown":
        return extract_txt(file_path)
    return ""

def extract_pdf(file_path: str):
    text = ""
    pdf = fitz.open(file_path)
    for page in pdf:
        text += page.get_text()
    pdf.close()
    return text

def extract_docx(file_path: str):
    document = Document(file_path)
    text = []
    for paragraph in document.paragraphs:
        text.append(paragraph.text)
    return "\n".join(text)

def extract_pptx(file_path: str):
    presentation = Presentation(file_path)
    text = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
    return "\n".join(text)

def extract_xlsx(file_path: str):
    workbook = load_workbook(file_path)
    text = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            values = []
            for cell in row:
                if cell is not None:
                    values.append(str(cell))
            text.append(" | ".join(values))
    return "\n".join(text)

def extract_txt(file_path: str):
    with open(file_path, "r", encoding="utf-8") as file:
        return file.read()

def extract_csv(file_path: str):
    text = []
    with open(file_path, newline="", encoding="utf-8") as file:
        reader = csv.reader(file)
        for row in reader:
            text.append(" | ".join(row))
    return "\n".join(text)

def extract_json(file_path: str):
    with open(file_path, "r", encoding="utf-8") as file:
        data = json.load(file)
    return json.dumps(
        data,
        indent=2,
        ensure_ascii=False
    )

def extract_html(file_path: str):
    with open(file_path, "r", encoding="utf-8") as file:
        html = file.read()
    soup = BeautifulSoup(
        html,
        "html.parser"
    )
    return soup.get_text(separator="\n")